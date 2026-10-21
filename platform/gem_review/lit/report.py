"""Step 3: "literature", "review_queue" and "logic_chains" sheets in review.xlsx."""

from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import EvidenceType, compartments_disagree

LITERATURE_COLUMNS = [
    "gene", "reaction_id", "reaction_name", "ec", "subsystem", "priority_tier", "pair_hits",
    "fulltext_available", "status", "best_evidence_level", "best_evidence_name",
    "no_organism_specific_evidence", "n_claims", "n_supporting", "claims_disagree", "pmids_cited", "quotes",
    "shared_with", "input_tokens", "output_tokens",
]
QUEUE_COLUMNS = ["priority_tier", "reason", "gene", "reaction_id", "reaction_name", "ec", "subsystem",
                 "pair_hits", "best_evidence_level", "n_claims", "n_supporting", "pmids_cited"]
CHAIN_COLUMNS = ["mode", "label", "gene", "reaction_id", "question", "status", "answer", "evidence_level",
                 "evidence_name", "verdict", "n_quotes", "n_inferences", "pmids_cited", "steps",
                 "input_tokens", "output_tokens", "requests", "error"]


def _disagree(claims: list[dict]) -> bool:
    verdicts = {c["supports_gene_reaction"] for c in claims}
    return len(verdicts) > 1 or compartments_disagree(c.get("compartment") for c in claims)


def literature_rows(retrievals: list[dict], results: list[dict]) -> list[dict]:
    by_pair = {(r["gene"], r["reaction_id"]): r for r in results}
    rows = []
    for rec in retrievals:
        res = by_pair.get((rec["gene"], rec["reaction_id"]))
        review = (res or {}).get("review") or {}
        claims = review.get("claims", [])
        level = review.get("best_evidence_level")
        rows.append({
            "gene": rec["gene"], "reaction_id": rec["reaction_id"], "reaction_name": rec["reaction_name"],
            "ec": ";".join(rec["ec_codes"]), "subsystem": rec["subsystem"], "priority_tier": rec.get("tier"),
            "pair_hits": rec["pair_hits"], "fulltext_available": "Y" if rec["fulltext_available"] else "N",
            "status": res["status"] if res else "not_extracted",
            "best_evidence_level": level,
            "best_evidence_name": EvidenceType(level).name if level is not None else None,
            "no_organism_specific_evidence": review.get("no_organism_specific_evidence"),
            "n_claims": len(claims), "n_supporting": sum(c["supports_gene_reaction"] for c in claims),
            "claims_disagree": _disagree(claims) if claims else False,
            "pmids_cited": ";".join(dict.fromkeys(c["pmid"] for c in claims)),
            "quotes": " | ".join(f'{c["pmid"]}: "{c["quote"]}"' for c in claims),
            "shared_with": (res or {}).get("shared_with"),
            "input_tokens": (res or {}).get("usage", {}).get("input_tokens"),
            "output_tokens": (res or {}).get("usage", {}).get("output_tokens"),
        })
    return rows


def queue_reason(row: dict) -> str | None:
    if row["status"] in {"EXTRACTION_FAILED", "MODEL_ERROR", "not_extracted"}:
        return row["status"]
    if row["claims_disagree"]:
        return "claims_disagree"
    if row["n_claims"] and not row["n_supporting"]:
        return "contradicted"   # every validated claim says the gene does NOT catalyse this reaction
    if row["status"] == "no_hits":
        return "no_hits"
    if row["best_evidence_level"] is not None and row["best_evidence_level"] <= EvidenceType.homology:
        return "low_evidence"
    return None


def review_queue_rows(rows: list[dict]) -> list[dict]:
    order = {"EXTRACTION_FAILED": 0, "MODEL_ERROR": 0, "not_extracted": 0, "contradicted": 1,
             "claims_disagree": 2, "low_evidence": 3, "no_hits": 4}
    queue = []
    for row in rows:
        reason = queue_reason(row)
        if reason:
            queue.append({**{k: row[k] for k in QUEUE_COLUMNS if k != "reason"}, "reason": reason})
    queue.sort(key=lambda r: (r["priority_tier"] if r["priority_tier"] is not None else 99,
                              order[r["reason"]], r["reaction_id"], r["gene"]))
    return queue


def _render_step(index: int, step: dict) -> str:
    if step["kind"] == "quote":
        level = step["evidence_type"] if step["organism_is_yarrowia"] else min(step["evidence_type"], 1)
        where = "" if step["organism_is_yarrowia"] else ", other organism"
        return f'[{index}] quote {step["pmid"]} {step["chunk_id"]} (L{level}{where}): "{step["quote"]}"'
    if step["kind"] == "inference":
        return f'[{index}] inference <- {step["rests_on"]}: {step["text"]}'
    return f'[{index}] verdict {step["answer"]} L{step["evidence_level"]} <- {step["rests_on"]}: {step["verdict"]}'


def chain_rows(chains: list[dict]) -> list[dict]:
    rows = []
    for rec in chains:
        steps = (rec.get("chain") or {}).get("steps", [])
        verdict = steps[-1] if steps else {}
        level = verdict.get("evidence_level")
        usage = rec.get("usage") or {}
        rows.append({
            "mode": rec["mode"], "label": rec["label"], "gene": rec.get("gene"),
            "reaction_id": rec.get("reaction_id"), "question": rec["question"], "status": rec["status"],
            "answer": verdict.get("answer"), "evidence_level": level,
            "evidence_name": EvidenceType(level).name if level is not None else None,
            "verdict": verdict.get("verdict"),
            "n_quotes": sum(s["kind"] == "quote" for s in steps),
            "n_inferences": sum(s["kind"] == "inference" for s in steps),
            "pmids_cited": ";".join(dict.fromkeys(s["pmid"] for s in steps if s["kind"] == "quote")),
            "steps": "\n".join(_render_step(i, s) for i, s in enumerate(steps)),
            "input_tokens": usage.get("input_tokens"), "output_tokens": usage.get("output_tokens"),
            "requests": usage.get("requests"), "error": rec.get("error"),
        })
    return rows


def write_review_xlsx(path: Path, rows: list[dict], queue: list[dict]) -> Path:
    """Add or replace the literature and review_queue sheets; other sheets are kept."""
    return write_sheets(path, [("literature", LITERATURE_COLUMNS, rows), ("review_queue", QUEUE_COLUMNS, queue)])


def write_chain_sheet(path: Path, rows: list[dict]) -> Path:
    """Add or replace the logic_chains sheet; other sheets are kept."""
    return write_sheets(path, [("logic_chains", CHAIN_COLUMNS, rows)])


def write_sheets(path: Path, sheets: list[tuple[str, list[str], list[dict]]]) -> Path:
    workbook = load_workbook(path) if path.exists() else Workbook()
    if not path.exists():
        workbook.remove(workbook.active)
    for name, columns, data in sheets:
        if name in workbook.sheetnames:
            workbook.remove(workbook[name])
        sheet = workbook.create_sheet(name)
        sheet.append(columns)
        for row in data:
            sheet.append([row.get(column) for column in columns])
        sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path

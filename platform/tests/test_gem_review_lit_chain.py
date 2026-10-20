"""Logic-chain layer with the agent mocked through pydantic-ai FunctionModel."""

import asyncio
import json
import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("PYDANTIC_AI_NO_BANNER", "1")

import httpx  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from pydantic import ValidationError  # noqa: E402
from pydantic_ai.messages import ModelResponse, ToolCallPart  # noqa: E402
from pydantic_ai.models.function import FunctionModel  # noqa: E402

from gem_review.lit.chain import (ChainContext, ChainResult, LogicChain, build_chain_agent,  # noqa: E402
                                  load_chains, pair_chain_inputs, question_chain_inputs,
                                  retrieve_question, run_chain, write_chain)
from gem_review.lit.client import EPMC_SEARCH_URL, CachedClient, DiskCache  # noqa: E402
from gem_review.lit.models import EvidenceType  # noqa: E402
from gem_review.lit.report import chain_rows, write_chain_sheet  # noqa: E402

CHUNK = ("Aox2 was shown to be a long-chain acyl-CoA oxidase and Aox3 was found to be active "
         "against short-chain fatty acids in Y. lipolytica.")
CHUNKS = {"111:abstract": {"chunk_id": "111:abstract", "pmid": "111", "source": "abstract", "text": CHUNK}}
CTX = ChainContext("Does POX3 oxidise C16 acyl-CoA?", {"111"}, CHUNKS)

QUOTE = {"kind": "quote", "pmid": "111", "chunk_id": "111:abstract",
         "quote": "Aox3 was found to be active against short-chain fatty acids",
         "evidence_type": 4, "organism_is_yarrowia": True}
INFER = {"kind": "inference", "text": "C16 is not a short-chain substrate, so the quote argues against C16.",
         "rests_on": [0]}
VERDICT = {"kind": "verdict", "answer": "no", "evidence_level": 4, "rests_on": [1],
           "verdict": "The chunks assign POX3 to short-chain substrates, not C16."}
VALID = [QUOTE, INFER, VERDICT]


def _scripted_model(answers: list[list[dict]]):
    """Each model call returns the next scripted step list; records every call."""
    calls = []

    def respond(messages, info):
        calls.append(messages)
        steps = answers[min(len(calls), len(answers)) - 1]
        return ModelResponse(parts=[ToolCallPart(info.output_tools[0].name, {"steps": steps})])

    return FunctionModel(respond), calls


def _run(model) -> ChainResult:
    agent = build_chain_agent(model, retries=2)
    result = ChainResult("question", "q_test", CTX.prompt, CTX.prompt, "pending")
    return asyncio.run(run_chain(agent, CTX, result, model_name="function"))


class StructureTest(unittest.TestCase):
    def test_valid_chain_passes(self):
        chain = LogicChain.model_validate({"steps": VALID})
        self.assertEqual(chain.verdict.evidence_level, EvidenceType.biochemical)
        self.assertEqual(len(chain.quotes), 1)

    def test_inference_without_a_quote_is_rejected(self):
        steps = [{"kind": "inference", "text": "POX3 is short-chain.", "rests_on": [0]}, VERDICT]
        with self.assertRaisesRegex(ValidationError, "must point to earlier steps"):
            LogicChain.model_validate({"steps": steps})
        # Resting only on an inference that itself does not trace fails too.
        steps = [QUOTE, {**INFER, "rests_on": [2]}, {**INFER, "rests_on": [1]}, {**VERDICT, "rests_on": [0]}]
        with self.assertRaisesRegex(ValidationError, "step 2: inference does not trace to any quote"):
            LogicChain.model_validate({"steps": steps})
        # Resting on a (misplaced) quote-less verdict does not trace.
        early = {**VERDICT, "answer": "unclear", "evidence_level": 0, "rests_on": []}
        with self.assertRaisesRegex(ValidationError, "step 1: inference does not trace to any quote"):
            LogicChain.model_validate({"steps": [early, {**INFER, "rests_on": [0]}, early]})
        with self.assertRaises(ValidationError):  # rests_on may not be empty
            LogicChain.model_validate({"steps": [QUOTE, {**INFER, "rests_on": []}, VERDICT]})

    def test_inference_tracing_is_transitive(self):
        steps = [QUOTE, INFER, {**INFER, "text": "Hence not C16.", "rests_on": [1]}, {**VERDICT, "rests_on": [2]}]
        self.assertEqual(LogicChain.model_validate({"steps": steps}).verdict.evidence_level, 4)

    def test_verdict_above_quote_level_is_rejected(self):
        weak = [{**QUOTE, "evidence_type": 2}, INFER, VERDICT]
        with self.assertRaisesRegex(ValidationError, "evidence_level 4 exceeds 2"):
            LogicChain.model_validate({"steps": weak})
        foreign = [{**QUOTE, "organism_is_yarrowia": False}, INFER, {**VERDICT, "evidence_level": 2}]
        with self.assertRaisesRegex(ValidationError, "evidence_level 2 exceeds 1"):
            LogicChain.model_validate({"steps": foreign})
        capped = [{**QUOTE, "organism_is_yarrowia": False}, INFER, {**VERDICT, "evidence_level": 1}]
        self.assertEqual(LogicChain.model_validate({"steps": capped}).verdict.evidence_level, 1)

    def test_chain_without_quotes_must_end_at_level_zero(self):
        with self.assertRaisesRegex(ValidationError, "exceeds 0"):
            LogicChain.model_validate({"steps": [{**VERDICT, "rests_on": []}]})
        ok = LogicChain.model_validate({"steps": [{**VERDICT, "answer": "unclear", "evidence_level": 0,
                                                    "rests_on": []}]})
        self.assertEqual(ok.verdict.evidence_level, EvidenceType.none)

    def test_verdict_must_be_last_and_unique(self):
        with self.assertRaisesRegex(ValidationError, "must end with a verdict"):
            LogicChain.model_validate({"steps": [QUOTE, INFER]})
        with self.assertRaisesRegex(ValidationError, "verdict must be the last step"):
            LogicChain.model_validate({"steps": [QUOTE, {**VERDICT, "rests_on": [0]}, VERDICT]})

    def test_quote_longer_than_25_words_is_rejected(self):
        with self.assertRaises(ValidationError):
            LogicChain.model_validate({"steps": [{**QUOTE, "quote": " ".join(["word"] * 26)}, VERDICT]})


class AgentRunTest(unittest.TestCase):
    def test_valid_chain_passes_in_one_call(self):
        model, calls = _scripted_model([VALID])
        result = _run(model)
        self.assertEqual(result.status, "ok")
        self.assertEqual(len(calls), 1)
        self.assertEqual(result.chain.verdict.answer, "no")
        # Chunks travel in the instructions, the question in the user message.
        request = calls[0][0]
        self.assertIn("111:abstract", request.instructions)
        user = [p for p in request.parts if p.part_kind == "user-prompt"][0].content
        self.assertEqual(user, CTX.prompt)

    def test_inference_without_quote_exhausts_retries_and_records_chain_failed(self):
        bad = [{"kind": "inference", "text": "POX3 is short-chain.", "rests_on": [0]},
               {**VERDICT, "evidence_level": 0, "rests_on": []}]
        model, calls = _scripted_model([bad])
        result = _run(model)
        self.assertEqual(result.status, "CHAIN_FAILED")
        self.assertEqual(len(calls), 3)  # first attempt + 2 retries
        self.assertIsNone(result.chain)
        self.assertIn("Exceeded maximum output retries", result.error)

    def test_verdict_above_quote_level_is_retried_then_fixed(self):
        over = [{**QUOTE, "organism_is_yarrowia": False}, INFER, VERDICT]
        fixed = [{**QUOTE, "organism_is_yarrowia": False}, INFER, {**VERDICT, "evidence_level": 1}]
        model, calls = _scripted_model([over, fixed])
        result = _run(model)
        self.assertEqual(result.status, "ok")
        self.assertEqual(len(calls), 2)
        self.assertIn("exceeds 1", str(calls[1][-1].parts[0].content))
        self.assertEqual(result.chain.verdict.evidence_level, EvidenceType.homology)

    def test_quote_not_verbatim_is_retried_like_validate_claims(self):
        paraphrase = [{**QUOTE, "quote": "POX3 only acts on short chains"}, INFER, VERDICT]
        wrong_pmid = [{**QUOTE, "pmid": "999"}, INFER, VERDICT]
        model, calls = _scripted_model([paraphrase, wrong_pmid, VALID])
        result = _run(model)
        self.assertEqual(result.status, "ok")
        self.assertEqual(len(calls), 3)
        self.assertIn("step 0: quote is not verbatim in chunk 111:abstract", str(calls[1][-1].parts[0].content))
        self.assertIn("step 0: pmid 999 is not among the retrieved PMIDs", str(calls[2][-1].parts[0].content))


class OutputTest(unittest.TestCase):
    def test_pair_chain_written_next_to_review_and_sheet_added(self):
        record = {"gene": "YALI1D32754g", "reaction_id": "R100", "reaction_name": "acyl-CoA oxidase (hexadecanoyl-CoA)",
                  "ec_codes": ["1.3.3.6"], "yali0": "YALI0D24750g", "symbols": ["POX3", "ACO3"],
                  "pmids": ["111"], "chunks": list(CHUNKS.values())}
        ctx, result = pair_chain_inputs(record)
        self.assertIn("YALI0D24750g, YALI0_D24750g, POX3, ACO3", ctx.prompt)
        model, _ = _scripted_model([VALID])
        result = asyncio.run(run_chain(build_chain_agent(model), ctx, result, model_name="function"))
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            path = write_chain(tmp / "lit", result)
            self.assertEqual(path.name, "YALI1D32754g__R100.chain.json")
            stored = json.loads(path.read_text())
            self.assertEqual(stored["chain"]["steps"][-1]["answer"], "no")
            book_path = tmp / "review.xlsx"
            write_chain_sheet(book_path, [{"mode": "x"}])  # replaced, not appended, on the next write
            rows = chain_rows(load_chains(tmp / "lit"))
            write_chain_sheet(book_path, rows)
            sheet = load_workbook(book_path)["logic_chains"]
            self.assertEqual(sheet.max_row, 2)
            header = [c.value for c in sheet[1]]
            row = dict(zip(header, [c.value for c in sheet[2]]))
            self.assertEqual((row["answer"], row["evidence_level"], row["evidence_name"]), ("no", 4, "biochemical"))
            self.assertIn('[0] quote 111 111:abstract (L4): "Aox3 was found', row["steps"])

    def test_question_retrieval_reports_zero_hits_without_model(self):
        seen = []

        def handler(request: httpx.Request) -> httpx.Response:
            seen.append(request.url.params.get("query"))
            assert str(request.url).split("?")[0] == EPMC_SEARCH_URL
            return httpx.Response(200, json={"hitCount": 0, "resultList": {"result": []}})

        with tempfile.TemporaryDirectory() as tmp:
            client = CachedClient(DiskCache(Path(tmp)), transport=httpx.MockTransport(handler), min_interval=0.0)

            async def go():
                try:
                    return await retrieve_question(client, "Does POX3 oxidise C16?", gene="YALI1D32754g",
                                                   aliases=["POX3", "YALI0D24750g"], terms=["acyl-CoA oxidase"])
                finally:
                    await client.aclose()

            record = asyncio.run(go())
        self.assertEqual((record["gene_hits"], record["pair_hits"], record["chunks"]), (0, 0, []))
        self.assertIn('"YALI1_D32754g" OR "POX3" OR "YALI0D24750g" OR "YALI0_D24750g"', seen[0])
        self.assertTrue(seen[1].endswith('AND ("acyl-CoA oxidase")'))
        ctx, result = question_chain_inputs(record, Path("x.retrieval.json"))
        self.assertIn("names this gene goes by in papers: YALI1D32754g, POX3, YALI0D24750g", ctx.prompt)
        self.assertEqual(result.label, record["slug"])


if __name__ == "__main__":
    unittest.main()

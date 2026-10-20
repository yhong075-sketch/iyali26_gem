"""Step-2 extraction with the agent mocked through pydantic-ai FunctionModel."""

import asyncio
import json
import os
import tempfile
import unittest
from pathlib import Path

os.environ.setdefault("PYDANTIC_AI_NO_BANNER", "1")

from openpyxl import load_workbook  # noqa: E402
from pydantic import ValidationError  # noqa: E402
from pydantic_ai import ModelRetry  # noqa: E402
from pydantic_ai.messages import ModelResponse, ToolCallPart  # noqa: E402
from pydantic_ai.models.function import FunctionModel  # noqa: E402

from gem_review.lit.extract import (PairContext, build_agent, load_results, run_extraction,  # noqa: E402
                                    validate_claims)
from gem_review.lit.models import Claim, ClaimSet, EvidenceType, PairReview  # noqa: E402
from gem_review.lit.report import literature_rows, review_queue_rows, write_review_xlsx  # noqa: E402

CHUNK_TEXT = "Deletion of POX2 (YALI0F10857g) abolished acyl-CoA oxidase activity on C18 substrates in Y. lipolytica."
CTX = PairContext("YALI1F14495g", "R1490", "acyl-CoA oxidase (icosanoyl-CoA)", ["1.3.3.6"], {"111", "222"},
                  {"PMC111:p1": {"chunk_id": "PMC111:p1", "pmid": "111", "source": "fulltext", "text": CHUNK_TEXT},
                   "222:abstract": {"chunk_id": "222:abstract", "pmid": "222", "source": "abstract",
                                    "text": "Lipid production in Yarrowia."}})
GOOD = {"pmid": "111", "supports_gene_reaction": True, "evidence_type": 3, "compartment": None,
        "organism_is_yarrowia": True, "quote": "Deletion of POX2 (YALI0F10857g) abolished acyl-CoA oxidase activity",
        "chunk_id": "PMC111:p1"}


def _claim(**overrides) -> Claim:
    return Claim(**{**GOOD, **overrides})


def _record(gene="YALI1F14495g", rid="R1490", hits=2, name="acyl-CoA oxidase (icosanoyl-CoA)"):
    return {"gene": gene, "reaction_id": rid, "reaction_name": name, "ec_codes": ["1.3.3.6"],
            "yali0": "YALI0F10857g", "symbols": ["POX2", "AOX2"],
            "subsystem": "Fatty acid degradation", "tier": 1, "pair_hits": hits, "pmids": ["111", "222"],
            "fulltext_available": True, "chunks": list(CTX.chunks.values()) if hits else []}


def _scripted_model(answers: list[list[dict]]):
    """Each model call returns the next scripted claim list; records how often it was called."""
    calls = []

    def respond(messages, info):
        calls.append(messages)
        claims = answers[min(len(calls), len(answers)) - 1]
        return ModelResponse(parts=[ToolCallPart(info.output_tools[0].name, {"claims": claims})])

    return FunctionModel(respond), calls


class ValidatorTest(unittest.TestCase):
    def test_pmid_not_in_retrieved_set_is_rejected(self):
        with self.assertRaises(ModelRetry) as caught:
            validate_claims(CTX, ClaimSet(claims=[_claim(pmid="999")]))
        self.assertIn("pmid 999 is not among the retrieved PMIDs", str(caught.exception))

    def test_quote_not_in_chunk_is_rejected(self):
        with self.assertRaises(ModelRetry) as caught:
            validate_claims(CTX, ClaimSet(claims=[_claim(quote="POX2 is the main oxidase for oleate")]))
        self.assertIn("not verbatim in chunk PMC111:p1", str(caught.exception))

    def test_chunk_must_belong_to_pmid(self):
        with self.assertRaises(ModelRetry):
            validate_claims(CTX, ClaimSet(claims=[_claim(pmid="222")]))
        with self.assertRaises(ModelRetry):
            validate_claims(CTX, ClaimSet(claims=[_claim(chunk_id="PMC999:p1")]))

    def test_valid_claim_passes_and_whitespace_is_normalised(self):
        out = validate_claims(CTX, ClaimSet(claims=[_claim(quote="Deletion of  POX2\n(YALI0F10857g) abolished")]))
        self.assertEqual(len(out.claims), 1)

    def test_quote_longer_than_25_words_is_rejected(self):
        with self.assertRaises(ValidationError):
            _claim(quote=" ".join(["word"] * 26))


class LevelTest(unittest.TestCase):
    def test_best_level_computed_in_python_and_capped_for_other_organisms(self):
        review = PairReview.from_claims("g", "r", [
            _claim(evidence_type=4, organism_is_yarrowia=False),      # capped to homology
            _claim(evidence_type=2, organism_is_yarrowia=True),
            _claim(evidence_type=4, supports_gene_reaction=False),    # contradiction never raises the level
        ])
        self.assertEqual(review.best_evidence_level, EvidenceType.physiological)
        self.assertFalse(review.no_organism_specific_evidence)
        self.assertTrue(review.claims_disagree())
        only_foreign = PairReview.from_claims("g", "r", [_claim(evidence_type=4, organism_is_yarrowia=False)])
        self.assertEqual(only_foreign.best_evidence_level, EvidenceType.homology)
        self.assertTrue(only_foreign.no_organism_specific_evidence)
        self.assertEqual(PairReview.from_claims("g", "r", []).best_evidence_level, EvidenceType.none)

    def test_compartment_spelling_variants_do_not_disagree(self):
        same = PairReview.from_claims("g", "r", [_claim(compartment="peroxisomal"),
                                                _claim(compartment="Peroxisomes"), _claim(compartment="peroxisome")])
        self.assertFalse(same.claims_disagree())
        different = PairReview.from_claims("g", "r", [_claim(compartment="cytosol"), _claim(compartment="mitochondrial")])
        self.assertTrue(different.claims_disagree())


class AgentRunTest(unittest.TestCase):
    def _run(self, model, records, tmp):
        agent = build_agent(model, retries=2)
        return asyncio.run(run_extraction(records, agent=agent, model_name="function", out_dir=tmp / "lit"))

    def test_bad_pmid_exhausts_retries_and_records_extraction_failed(self):
        model, calls = _scripted_model([[{**GOOD, "pmid": "999"}]])
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            results = self._run(model, [_record()], tmp)
            self.assertEqual(results[0].status, "EXTRACTION_FAILED")
            self.assertEqual(len(calls), 3)  # first attempt + 2 retries
            self.assertIsNone(results[0].review)
            stored = json.loads((tmp / "lit" / "YALI1F14495g__R1490.review.json").read_text())
            self.assertEqual(stored["status"], "EXTRACTION_FAILED")
            self.assertIn("Exceeded maximum output retries", stored["error"])

    def test_retry_then_valid_answer_is_kept(self):
        model, calls = _scripted_model([[{**GOOD, "quote": "not in the chunk at all"}], [GOOD]])
        with tempfile.TemporaryDirectory() as tmp:
            results = self._run(model, [_record()], Path(tmp))
            self.assertEqual(results[0].status, "ok")
            self.assertEqual(len(calls), 2)
            # The question carries the aliases papers actually use.
            question = [p for p in calls[0][0].parts if p.part_kind == "user-prompt"][0].content
            self.assertIn("YALI1_F14495g, YALI0F10857g, YALI0_F10857g, POX2, AOX2", question)
            self.assertEqual(results[0].review.best_evidence_level, EvidenceType.genetic)
            # The retry prompt carried the validator's message back to the model.
            self.assertIn("not verbatim", str(calls[1][-1].parts[0].content))

    def test_no_hits_pairs_skip_the_model_and_identical_questions_share_a_call(self):
        model, calls = _scripted_model([[GOOD]])
        records = [_record(hits=0, rid="R0"), _record(rid="R1490"), _record(rid="R1491")]
        with tempfile.TemporaryDirectory() as tmp:
            tmp = Path(tmp)
            results = self._run(model, records, tmp)
            by_rid = {r.reaction_id: r for r in results}
            self.assertEqual(by_rid["R0"].status, "no_hits")
            self.assertEqual(by_rid["R0"].review.best_evidence_level, EvidenceType.none)
            self.assertEqual(len(calls), 1)
            self.assertEqual(by_rid["R1491"].shared_with, "YALI1F14495g__R1490")
            self.assertEqual(by_rid["R1491"].review.claims, by_rid["R1490"].review.claims)
            usage_lines = (tmp / "lit" / "usage.jsonl").read_text().splitlines()
            self.assertEqual(len(usage_lines), 1)

            rows = literature_rows(records, load_results(tmp / "lit"))
            queue = review_queue_rows(rows)
            self.assertEqual([r["reaction_id"] for r in queue], ["R0"])  # genetic-level pairs are not queued
            contradicted = {**rows[1], "n_claims": 2, "n_supporting": 0, "best_evidence_level": 0}
            self.assertEqual(review_queue_rows([contradicted])[0]["reason"], "contradicted")
            xlsx = write_review_xlsx(tmp / "review.xlsx", rows, queue)
            book = load_workbook(xlsx)
            self.assertEqual(book.sheetnames, ["literature", "review_queue"])
            self.assertEqual(book["literature"].max_row, 4)
            self.assertEqual(book["review_queue"]["B2"].value, "no_hits")


if __name__ == "__main__":
    unittest.main()
